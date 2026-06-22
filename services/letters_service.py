"""
services/letter_generation_service.py

Business logic for the HR Letter Generation module:
  - request code generation (LTR-REQ-2024-001)
  - template placeholder rendering
  - AI (OpenAI) letter body generation
  - PDF / Excel export helpers
"""

import io
import re
from datetime import date, datetime
from typing import Optional

from sqlalchemy.orm import Session
from sqlalchemy import func

from core.config import settings
from model.HR_Operations.letter_generation import LetterRequest, LetterTemplate
from model.onboarding.employee import Employee


# ======================================================
# REQUEST CODE GENERATOR -> "LTR-REQ-2024-001"
# ======================================================
def generate_request_code(db: Session) -> str:
    year = datetime.utcnow().year
    prefix = f"LTR-REQ-{year}-"

    last = (
        db.query(LetterRequest)
        .filter(LetterRequest.request_code.like(f"{prefix}%"))
        .order_by(LetterRequest.id.desc())
        .first()
    )

    if last:
        try:
            last_seq = int(last.request_code.split("-")[-1])
        except (ValueError, IndexError):
            last_seq = 0
    else:
        last_seq = 0

    return f"{prefix}{last_seq + 1:03d}"


# ======================================================
# PLACEHOLDER RENDERING -> fills {{employee_name}} etc.
# ======================================================
def build_placeholder_context(employee: Employee, extra: Optional[dict] = None) -> dict:
    full_name = " ".join(
        p for p in [employee.first_name, employee.middle_name, employee.last_name] if p
    )

    context = {
        "employee_name": full_name,
        "employee_code": employee.employee_code,
        "department": employee.department or "",
        "designation": employee.designation or "",
        "location": employee.location or "",
        "joining_date": employee.joining_date.strftime("%d-%b-%Y") if employee.joining_date else "",
        "today": date.today().strftime("%d-%b-%Y"),
        "company_name": "Levitica Technologies",
    }

    if extra:
        context.update({k: str(v) for k, v in extra.items()})

    return context


def render_template_string(template_str: str, context: dict) -> str:
    """Replaces {{key}} tokens in a template string with values from context."""

    def _replace(match: re.Match) -> str:
        key = match.group(1).strip()
        return str(context.get(key, match.group(0)))

    return re.sub(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}", _replace, template_str)


# ======================================================
# AI GENERATION (OpenAI) -> powers the "AI" button
# ======================================================
_ai_client = None


def _get_ai_client():
    """Lazily create the OpenAI client so the app still boots without a key configured."""
    global _ai_client
    if _ai_client is None:
        from openai import OpenAI
        _ai_client = OpenAI(api_key=settings.OPENAI_API_KEY)
    return _ai_client


def ai_generate_letter(
    letter_type: str,
    context: dict,
    tone: str = "formal",
    extra_instructions: Optional[str] = None,
) -> dict:
    """
    Calls OpenAI to draft a subject + body for an HR letter.
    Returns {"subject": str, "body": str}
    """
    client = _get_ai_client()

    details_lines = "\n".join(f"- {k}: {v}" for k, v in context.items() if v)

    prompt = f"""
You are an HR executive drafting an official company letter.

Letter type: {letter_type}
Tone: {tone}

Employee / letter details:
{details_lines}

{f"Additional instructions: {extra_instructions}" if extra_instructions else ""}

Write the letter as JSON with exactly two keys:
- "subject": a short professional subject line for the letter
- "body": the full letter body, addressed properly, with company name,
  date, salutation, clear paragraphs, and a closing signature block.
  Do not use markdown formatting, just plain text suitable for a printed letter.

Return only valid JSON, nothing else.
"""

    resp = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[
            {"role": "system", "content": "You write precise, professional HR letters and return JSON only."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.4,
    )

    text = resp.choices[0].message.content.strip()

    import json
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        text = text[text.find("{"): text.rfind("}") + 1]
        data = json.loads(text)

    return {
        "subject": data.get("subject", f"{letter_type}"),
        "body": data.get("body", ""),
    }


# ======================================================
# DASHBOARD STATS
# ======================================================
def get_dashboard_stats(db: Session) -> dict:
    total_templates = db.query(func.count(LetterTemplate.id)).scalar() or 0
    ai_optimized = (
        db.query(func.count(LetterTemplate.id))
        .filter(LetterTemplate.is_ai_optimized.is_(True))
        .scalar()
        or 0
    )

    total_requests = db.query(func.count(LetterRequest.id)).scalar() or 0
    approved = (
        db.query(func.count(LetterRequest.id))
        .filter(LetterRequest.status == "APPROVED")
        .scalar()
        or 0
    )
    pending = (
        db.query(func.count(LetterRequest.id))
        .filter(LetterRequest.status == "PENDING")
        .scalar()
        or 0
    )
    rejected = (
        db.query(func.count(LetterRequest.id))
        .filter(LetterRequest.status == "REJECTED")
        .scalar()
        or 0
    )
    auto_approved = (
        db.query(func.count(LetterRequest.id))
        .filter(LetterRequest.auto_approved.is_(True))
        .scalar()
        or 0
    )
    total_downloads = db.query(func.sum(LetterRequest.download_count)).scalar() or 0

    category_rows = (
        db.query(LetterTemplate.category, func.count(LetterTemplate.id))
        .group_by(LetterTemplate.category)
        .all()
    )

    return {
        "total_templates": total_templates,
        "ai_optimized_templates": ai_optimized,
        "total_requests": total_requests,
        "approved_requests": approved,
        "pending_requests": pending,
        "rejected_requests": rejected,
        "auto_approved_requests": auto_approved,
        "total_downloads": int(total_downloads),
        "category_breakdown": [{"category": c, "count": n} for c, n in category_rows],
    }


# ======================================================
# EXPORT HELPERS
# ======================================================
def export_requests_to_excel(rows: list) -> io.BytesIO:
    """
    rows: list of (LetterRequest, Employee) tuples, e.g. from
          db.execute(select(LetterRequest, Employee).join(...)).all()
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Letter Requests"

    headers = ["Request ID", "Type", "Employee", "Department", "Status", "Date"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    for r, emp in rows:
        full_name = ""
        dept = ""
        if emp:
            full_name = " ".join(p for p in [emp.first_name, emp.middle_name, emp.last_name] if p)
            dept = emp.department or ""

        ws.append([
            r.request_code,
            r.letter_type,
            full_name,
            dept,
            r.status,
            r.letter_date.strftime("%d-%m-%Y") if r.letter_date else "",
        ])

    for col in ws.columns:
        max_len = max(len(str(c.value)) if c.value else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_letter_to_pdf(letter_request: LetterRequest, employee: Employee) -> io.BytesIO:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=25 * mm,
        bottomMargin=25 * mm,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "LetterTitle", parent=styles["Heading2"], spaceAfter=14
    )
    body_style = ParagraphStyle(
        "LetterBody", parent=styles["Normal"], fontSize=11, leading=16
    )
    meta_style = ParagraphStyle(
        "LetterMeta", parent=styles["Normal"], fontSize=9, textColor="#666666"
    )

    elements = []
    elements.append(Paragraph("Levitica Technologies", title_style))
    elements.append(
        Paragraph(
            f"Date: {letter_request.letter_date.strftime('%d-%b-%Y')} &nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Ref: {letter_request.request_code}",
            meta_style,
        )
    )
    elements.append(Spacer(1, 14))
    elements.append(Paragraph(f"<b>{letter_request.subject}</b>", styles["Heading3"]))
    elements.append(Spacer(1, 10))

    for para in letter_request.body.split("\n"):
        if para.strip():
            elements.append(Paragraph(para.strip(), body_style))
            elements.append(Spacer(1, 8))

    doc.build(elements)
    buffer.seek(0)
    return buffer